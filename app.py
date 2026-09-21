"""
BOB Reputation Risk Analyzer — Streamlit Dashboard
====================================================
A media-sentiment / reputation-risk analyzer that scrapes recent news
coverage for one or more companies, scores it with three independent
sentiment engines (a lexicon-based Net Percentage Score, NLTK VADER,
and a transformer-based RoBERTa sentiment model), and renders an
interactive dashboard with downloadable Excel reports.

Converted from a Flask + ngrok / Colab notebook into a single,
self-contained Streamlit app.
"""

import ast
import io
import string
import time
import zipfile
import warnings
from datetime import datetime, date

import numpy as np
import pandas as pd
import requests
import streamlit as st
from bs4 import BeautifulSoup

import re
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

import base64

from urllib.parse import (
    quote_plus,
    unquote,
    urlparse,
    parse_qs,
    urljoin
)

from google import genai

GEMINI_API_KEY = st.secrets["GEMINI_API_KEY"]

warnings.filterwarnings("ignore")

# ----------------------------------------------------------------------------
# Page config
# ----------------------------------------------------------------------------
st.set_page_config(
    page_title="BOB Reputation Risk Analyzer",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ----------------------------------------------------------------------------
# BOB brand theme (orange / maroon) — dashboard styling
# ----------------------------------------------------------------------------
BOB_ORANGE = "#F7931E"
BOB_MAROON = "#6E1E33"
BOB_GOLD = "#C89B3C"
BOB_CREAM = "#FFF8F0"
BOB_GREEN = "#1FA84C"
BOB_AMBER = "#F4A300"
BOB_RED = "#D93A3A"

CUSTOM_CSS = f"""
<style>
    .stApp {{
        background: linear-gradient(180deg, {BOB_CREAM} 0%, #ffffff 320px);
    }}
    .bob-banner {{
        background: linear-gradient(90deg, {BOB_MAROON} 0%, {BOB_ORANGE} 100%);
        padding: 1.6rem 2rem;
        border-radius: 14px;
        margin-bottom: 1.4rem;
        box-shadow: 0 6px 18px rgba(110, 30, 51, 0.25);
    }}
    .bob-banner h1 {{
        color: #ffffff;
        margin: 0;
        font-size: 1.9rem;
        font-weight: 800;
        letter-spacing: 0.3px;
    }}
    .bob-banner p {{
        color: #ffe9d1;
        margin: 0.3rem 0 0 0;
        font-size: 0.98rem;
    }}
    .kpi-card {{
        border-radius: 12px;
        padding: 1rem 1.1rem;
        color: white;
        text-align: center;
        box-shadow: 0 3px 10px rgba(0,0,0,0.12);
    }}
    .kpi-label {{
        font-size: 0.82rem;
        opacity: 0.9;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }}
    .kpi-value {{
        font-size: 2rem;
        font-weight: 800;
        margin-top: 0.15rem;
    }}
    .section-title {{
        color: {BOB_MAROON};
        font-weight: 800;
        border-bottom: 3px solid {BOB_ORANGE};
        display: inline-block;
        padding-bottom: 0.2rem;
        margin-top: 1.4rem;
        margin-bottom: 0.8rem;
    }}
    .stButton>button {{
        background: {BOB_MAROON};
        color: white;
        border-radius: 8px;
        border: none;
        font-weight: 600;
    }}
    .stButton>button:hover {{
        background: {BOB_ORANGE};
        color: white;
    }}
    div[data-testid="stDownloadButton"] button {{
        background: {BOB_ORANGE};
        color: white;
        border-radius: 8px;
        border: none;
        font-weight: 700;
    }}
    div[data-testid="stDownloadButton"] button:hover {{
        background: {BOB_MAROON};
    }}
    section[data-testid="stSidebar"] {{
        background: #fffaf3;
        border-right: 2px solid #f0dcc0;
    }}
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/114.0.0.0 Safari/537.36"
    )
}

# Yahoo's search backend frequently returns a bare "500 INKApi Error" to
# requests that look too much like a bot (missing consent cookies, a bare
# Accept header, no Referer). A fuller browser-like header set plus a
# same-session warm-up hit on the homepage (see Tier 5 below) reduces this.
YAHOO_HEADERS = {
    **HEADERS,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://search.yahoo.com/",
}

# ----------------------------------------------------------------------------
# Cached resources
# ----------------------------------------------------------------------------
@st.cache_resource(show_spinner="Loading sentiment engines (first run only)...")
def load_engines():
    """Load NLTK VADER and the transformer sentiment model once per session."""
    import nltk
    from nltk.sentiment import SentimentIntensityAnalyzer
    from transformers import (
        AutoModelForSequenceClassification,
        AutoTokenizer,
        AutoConfig,
        logging as hf_logging,
    )

    hf_logging.set_verbosity_error()
    nltk.download("vader_lexicon", quiet=True)

    vader = SentimentIntensityAnalyzer()

    model_name = "cardiffnlp/twitter-roberta-base-sentiment-latest"
    tokenizer = AutoTokenizer.from_pretrained(model_name)
    config = AutoConfig.from_pretrained(model_name)
    model = AutoModelForSequenceClassification.from_pretrained(model_name)

    return vader, tokenizer, config, model


def text_preprocess(text: str) -> str:
    """Normalize @mentions and links, matching the original preprocessing."""
    new_text = []
    for t in text.split(" "):
        t = "@user" if t.startswith("@") and len(t) > 1 else t
        t = "http" if t.startswith("http") else t
        new_text.append(t)
    return " ".join(new_text)


def safe_key(text: str) -> str:
    """Turn arbitrary text (e.g. a company name) into a safe Streamlit widget key fragment."""
    return "".join(c if c.isalnum() else "_" for c in text)


# ----------------------------------------------------------------------------
# Retrying HTTP session — shared by every tier below and by fetch_articles.
# Automatically retries transient failures (timeouts, connection resets,
# 429/500/502/503/504) with exponential backoff, which on its own removes a
# large share of "no articles found" cases caused by a single flaky request.
# ----------------------------------------------------------------------------
def _build_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=3,
        backoff_factor=0.6,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET", "POST", "HEAD"],
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


SESSION = _build_session()



# ----------------------------------------------------------------------------
# News discovery
#
# 8-tier fallback:
#
#   1. Google
#   2. googlesearch
#   3. DuckDuckGo via ddgs
#   4. Bing
#   5. Yahoo
#   6. DuckDuckGo HTML
#   7. Gemini Flash Lite
#   8. Direct Moneycontrol URL construction
#
# Each tier only runs when previous tiers found nothing.
# Every returned URL is strictly validated as a genuine
# Moneycontrol /news/tags/ URL.
# ----------------------------------------------------------------------------

def find_news_links(company_search_text: str, log=lambda m: None) -> list:

    links = []

    query_text = (
        f"{company_search_text} moneycontrol latest and breaking news"
    )

    # ========================================================================
    # HELPER 1: Validate Moneycontrol tag URL
    # ========================================================================

    def is_moneycontrol_tag_url(url: str) -> bool:

        if not isinstance(url, str):
            return False

        url = url.strip()

        if not url.startswith(("http://", "https://")):
            return False

        try:
            parsed = urlparse(url)

            hostname = (
                parsed.hostname or ""
            ).lower().rstrip(".")

            path = (
                parsed.path or ""
            ).lower()

            # Only genuine Moneycontrol domain
            if hostname not in {
                "www.moneycontrol.com",
                "moneycontrol.com",
            }:
                return False

            # Only news tag pages
            if not path.startswith("/news/tags/"):
                return False

            # Must contain a tag slug
            slug = path[len("/news/tags/"):].strip("/")

            if not slug:
                return False

            return True

        except Exception:
            return False

    # ========================================================================
    # HELPER 2: Normalize Moneycontrol URL
    # ========================================================================

    def normalize_moneycontrol_url(url: str):

        if not isinstance(url, str):
            return None

        try:

            url = unquote(
                url.strip()
            )

            if not is_moneycontrol_tag_url(url):
                return None

            parsed = urlparse(url)

            clean_url = (
                "https://www.moneycontrol.com"
                + parsed.path
            )

            if is_moneycontrol_tag_url(clean_url):
                return clean_url

        except Exception:
            pass

        return None

    # ========================================================================
    # HELPER 3: Add only valid URL
    # ========================================================================

    def add_link(url: str) -> bool:

        clean_url = normalize_moneycontrol_url(url)

        if (
            clean_url
            and clean_url not in links
        ):
            links.append(clean_url)
            return True

        return False

    # ========================================================================
    # TIER 1: GOOGLE
    # ========================================================================

    try:

        log(
            "Searching Google for Moneycontrol news links..."
        )

        url = (
            "https://www.google.com/search?q="
            + quote_plus(
                f'site:moneycontrol.com/news/tags/ '
                f'"{company_search_text}"'
            )
        )

        source = SESSION.get(
            url,
            headers=HEADERS,
            timeout=15
        )

        source.raise_for_status()

        soup = BeautifulSoup(
            source.text,
            "html.parser"
        )

        for a in soup.find_all(
            "a",
            href=True
        ):

            href = a["href"].strip()

            add_link(href)

    except Exception as e:

        log(
            f"Google scrape failed: {e}"
        )

    # ========================================================================
    # TIER 2: GOOGLESEARCH
    # ========================================================================

    if not links:

        try:

            log(
                "Trying secondary search path (googlesearch)..."
            )

            from googlesearch import search as g_search

            wrong_count = 0

            for link in g_search(
                f"{company_search_text} moneycontrol "
                f"latest and breaking news"
            ):

                if add_link(link):

                    wrong_count = 0

                else:

                    wrong_count += 1

                if wrong_count >= 10:
                    break

        except Exception as e:

            log(
                f"googlesearch fallback failed: {e}"
            )

    # ========================================================================
    # TIER 3: DUCKDUCKGO - DDGS
    # ========================================================================

    if not links:

        try:

            log(
                "Trying tertiary search path "
                "(DuckDuckGo via ddgs)..."
            )

            from ddgs import DDGS

            with DDGS() as ddgs:

                results = ddgs.text(
                    f"{company_search_text} "
                    f"latest news Moneycontrol",
                    max_results=30
                )

                for res in results:

                    href = res.get(
                        "href",
                        ""
                    )

                    if add_link(href):

                        if len(links) >= 3:
                            break

        except Exception as e:

            log(
                f"DuckDuckGo (ddgs) fallback failed: {e}"
            )

    # ========================================================================
    # TIER 4: BING
    # ========================================================================

    if not links:

        try:

            log(
                "Trying Bing search scrape..."
            )

            url = (
                "https://www.bing.com/search?q="
                + quote_plus(
                    f'site:moneycontrol.com/news/tags/ '
                    f'"{company_search_text}"'
                )
            )

            source = SESSION.get(
                url,
                headers=HEADERS,
                timeout=15
            )

            source.raise_for_status()

            soup = BeautifulSoup(
                source.text,
                "html.parser"
            )

            for a in soup.find_all(
                "a",
                href=True
            ):

                href = a["href"].strip()

                # ------------------------------------------------------------
                # Direct Moneycontrol URL
                # ------------------------------------------------------------

                if add_link(href):
                    continue

                # ------------------------------------------------------------
                # Bing redirect URL
                # ------------------------------------------------------------

                if "bing.com/ck/a" not in href.lower():
                    continue

                try:

                    parsed = urlparse(href)

                    params = parse_qs(
                        parsed.query
                    )

                    encoded_url = params.get(
                        "u",
                        [None]
                    )[0]

                    if not encoded_url:
                        continue

                    encoded_url = encoded_url.strip()

                    # Bing commonly prefixes encoded URL with a1
                    if encoded_url.startswith("a1"):
                        encoded_url = encoded_url[2:]

                    # Base64 padding
                    encoded_url += "=" * (
                        (-len(encoded_url)) % 4
                    )

                    try:

                        decoded_url = (
                            base64.urlsafe_b64decode(
                                encoded_url
                            )
                            .decode(
                                "utf-8",
                                errors="ignore"
                            )
                            .strip()
                        )

                    except Exception:

                        # Ignore unrelated Bing tracking URLs.
                        continue

                    decoded_url = unquote(
                        decoded_url
                    ).strip()

                    # IMPORTANT:
                    #
                    # add_link() parses hostname/path.
                    #
                    # Therefore a URL such as:
                    #
                    # https://www.bing.com/copilotsearch?
                    # q=site:moneycontrol.com/news/tags/
                    #
                    # will NOT pass validation.

                    add_link(
                        decoded_url
                    )

                except Exception:

                    continue

            log(
                f"Bing found {len(links)} "
                f"valid Moneycontrol tag URL(s)"
            )

            for link in links:

                log(
                    f"  -> {link}"
                )

        except Exception as e:

            log(
                f"Bing scrape failed: {e}"
            )

    # ========================================================================
    # TIER 5: YAHOO
    # ========================================================================

    if not links:

        try:

            log(
                "Trying Yahoo search scrape..."
            )

            # Warm up Yahoo session
            try:

                SESSION.get(
                    "https://search.yahoo.com/",
                    headers=YAHOO_HEADERS,
                    timeout=10
                )

            except Exception:
                pass

            url = (
                "https://search.yahoo.com/search?p="
                + quote_plus(query_text)
            )

            source = SESSION.get(
                url,
                headers=YAHOO_HEADERS,
                timeout=15
            )

            source.raise_for_status()

            soup = BeautifulSoup(
                source.text,
                "html.parser"
            )

            for a in soup.find_all(
                "a",
                href=True
            ):

                href = a["href"].strip()

                # Yahoo redirect format
                if "RU=" in href:

                    try:

                        href = unquote(
                            href
                            .split("RU=", 1)[1]
                            .split("/RK=", 1)[0]
                        )

                    except Exception:

                        continue

                add_link(href)

        except Exception as e:

            log(
                f"Yahoo scrape failed: {e}"
            )

    # ========================================================================
    # TIER 6: DUCKDUCKGO HTML
    # ========================================================================

    if not links:

        try:

            log(
                "Trying DuckDuckGo HTML-lite scrape..."
            )

            resp = SESSION.post(
                "https://html.duckduckgo.com/html/",
                data={
                    "q": (
                        f"{company_search_text} "
                        f"moneycontrol latest news"
                    )
                },
                headers=HEADERS,
                timeout=15
            )

            resp.raise_for_status()

            soup = BeautifulSoup(
                resp.text,
                "html.parser"
            )

            for a in soup.find_all(
                "a",
                href=True
            ):

                href = a["href"].strip()

                if add_link(href):

                    if len(links) >= 3:
                        break

        except Exception as e:

            log(
                f"DuckDuckGo HTML-lite scrape failed: {e}"
            )

    # ========================================================================
    # TIER 7: GEMINI
    #
    # IMPORTANT:
    # This is deliberately at the SAME indentation level as Tier 6.
    # It is NOT inside Tier 6.
    # ========================================================================

    if not links:

        try:

            log(
                "Trying Gemini Moneycontrol link finder..."
            )

            from google import genai

            # ================================================================
            # BACKEND API KEY
            #
            # Keep your API key in your backend.
            #
            # Expected existing variable:
            #
            # GEMINI_API_KEY
            #
            # No frontend input is required.
            # ================================================================

            gemini_client = genai.Client(
                api_key=GEMINI_API_KEY
            )

            prompt = f"""
Find the official Moneycontrol news TAG page for this company.

Company:
{company_search_text}

I need the actual Moneycontrol company/news tag page.

The URL MUST be on:
www.moneycontrol.com

The URL MUST use this path:
 /news/tags/

Expected format:
https://www.moneycontrol.com/news/tags/<company-slug>.html

Do NOT return:
- Google search URLs
- Bing search URLs
- Yahoo search URLs
- DuckDuckGo URLs
- Moneycontrol search URLs
- Moneycontrol Copilot URLs
- Moneycontrol stock/quote pages
- individual news article URLs
- any other website

Return only valid Moneycontrol tag URLs.
Return one URL per line.
Do not use markdown.
Do not use code fences.
Do not provide explanations.

If there are multiple valid tag URLs for this company,
return each URL on a separate line.
"""

            response = gemini_client.models.generate_content(
                model="gemini-flash-lite-latest",
                contents=prompt
            )

            gemini_text = (
                getattr(
                    response,
                    "text",
                    ""
                )
                or ""
            ).strip()

            if gemini_text:

                candidates = re.findall(
                    r'https?://[^\s<>"\']+',
                    gemini_text
                )

                for candidate in candidates:

                    candidate = candidate.rstrip(
                        ".,;:)]}>\"'"
                    )

                    if add_link(candidate):

                        log(
                            "Gemini found Moneycontrol URL: "
                            f"{candidate}"
                        )

            log(
                f"Gemini found {len(links)} "
                f"valid Moneycontrol tag URL(s)"
            )

            for link in links:

                log(
                    f"  -> {link}"
                )

        except Exception as e:

            log(
                f"Gemini link finder failed: {e}"
            )

    # ========================================================================
    # TIER 8: DIRECT MONEYCONTROL URL CONSTRUCTION
    #
    # Final fallback — no search engine and no AI.
    # ========================================================================

    if not links:

        try:

            log(
                "Trying direct Moneycontrol tag URL construction..."
            )

            base = (
                company_search_text
                .lower()
                .strip()
            )

            slug_variants = [
                re.sub(
                    r"[^a-z0-9]+",
                    "-",
                    base
                ).strip("-"),

                re.sub(
                    r"[^a-z0-9]+",
                    "",
                    base
                )
            ]

            slug_variants = list(
                dict.fromkeys(
                    slug
                    for slug in slug_variants
                    if slug
                )
            )

            for slug in slug_variants:

                candidate_url = (
                    "https://www.moneycontrol.com"
                    f"/news/tags/{slug}.html"
                )

                try:

                    response = SESSION.get(
                        candidate_url,
                        headers=HEADERS,
                        timeout=15,
                        allow_redirects=True
                    )

                    if response.status_code != 200:
                        continue

                    # Check the final URL after redirects.
                    final_url = response.url

                    if add_link(final_url):

                        log(
                            "Direct fallback found: "
                            f"{final_url}"
                        )

                    elif add_link(candidate_url):

                        log(
                            "Direct fallback found: "
                            f"{candidate_url}"
                        )

                except Exception:

                    continue

        except Exception as e:

            log(
                f"Direct URL construction failed: {e}"
            )

    # ========================================================================
    # FINAL SAFETY VALIDATION
    # ========================================================================

    valid_links = []

    for link in links:

        clean_url = normalize_moneycontrol_url(
            link
        )

        if (
            clean_url
            and clean_url not in valid_links
        ):
            valid_links.append(
                clean_url
            )

    links = valid_links

    # ========================================================================
    # FINAL LOG
    # ========================================================================

    log(
        f"Final valid Moneycontrol source pages: "
        f"{len(links)}"
    )

    for link in links:

        log(
            f"  -> {link}"
        )

    return links
    

def fetch_articles(links: list, start_date: datetime, end_date: datetime, log=lambda m: None) -> pd.DataFrame:
    articles = []
    for base_url in links:
        counter = 0
        time.sleep(np.random.randint(1, 3))
        log(f"Reading: {base_url}")
        query_params = {"page": 1}
        stop_company = False

        while counter < 100:
            try:
                response = SESSION.get(base_url, params=query_params, headers=HEADERS, timeout=15)
            except Exception:
                break
            if response.status_code != 200:
                break

            soup = BeautifulSoup(response.text, "html.parser")
            article_elements = soup.find_all("li", class_="clearfix")
            if not article_elements:
                break

            break_counter = 0
            for el in article_elements:
                try:
                    title = el.find("h2").text.strip()
                except Exception:
                    title = ""
                    break_counter += 1
                try:
                    link = el.find("a")["href"]
                except Exception:
                    link = ""
                    break_counter += 1
                try:
                    span_start = str(el).find("<span>")
                    span_end = str(el).find("</span>")
                    date_str = str(el)[span_start + 6 : span_end]
                except Exception:
                    date_str = ""
                    break_counter += 1
                try:
                    summary = el.find("p").text.strip()
                except Exception:
                    summary = ""
                    break_counter += 1

                try:
                    trimmed = date_str[: len(date_str) - 4]
                    article_date = datetime.strptime(trimmed, "%B %d, %Y %I:%M %p")
                except Exception:
                    break_counter += 1
                    continue

                if article_date < start_date:
                    stop_company = True
                    break

                if start_date <= article_date <= end_date:
                    articles.append(
                        {"title": title, "link": link, "date": article_date, "summary": summary}
                    )

            query_params["page"] += 1
            counter += 1
            if break_counter >= 28 or stop_company:
                break

    return pd.DataFrame(articles)


# ----------------------------------------------------------------------------
# Sentiment scoring
# ----------------------------------------------------------------------------
def score_sentiment(df_articles: pd.DataFrame, custom_words: dict, vader, tokenizer, config, model) -> pd.DataFrame:
    lexicon = vader.lexicon
    base_custom = {
        "default": -4, "defaulted": -4, "defaults": -4, "defaulting": -4,
        "defaulter": -4, "defaulters": -4, "up": 3, "down": -4,
    }
    base_custom.update(custom_words)
    lexicon.update(base_custom)

    df = df_articles.rename(columns={"link": "article_url", "title": "article_title",
                                      "date": "article_date", "summary": "article_summary"})
    df = df.drop_duplicates(subset=["article_title", "article_summary"]).reset_index(drop=True)
    if df.empty:
        return df

    neg, neu, pos, comp = [], [], [], []
    pos_words_list, neg_words_list = [], []
    for i in range(len(df)):
        text = f"{df.loc[i, 'article_title']}. {df.loc[i, 'article_summary']}"
        scores = vader.polarity_scores(text)
        neg.append(scores["neg"]); neu.append(scores["neu"])
        pos.append(scores["pos"]); comp.append(scores["compound"])

        p_words, n_words = [], []
        for word in text.split():
            s = vader.polarity_scores(word)["compound"]
            if s > 0:
                p_words.append(word)
            elif s < 0:
                n_words.append(word)
        pos_words_list.append(p_words)
        neg_words_list.append(n_words)

    df["negative"] = neg
    df["neutral"] = neu
    df["positive"] = pos
    df["compound"] = comp
    df["positive_words_list"] = pos_words_list
    df["negative_words_list"] = neg_words_list

    denom = (df["positive_words_list"].str.len() + df["negative_words_list"].str.len())
    df["NPS"] = ((df["positive_words_list"].str.len() - df["negative_words_list"].str.len()) / denom).round(2)
    df["NPS"] = df["NPS"].fillna(0)

    llm_sentiment = []
    import torch
    
    progress = st.progress(0, text="Starting sentiment scoring...")
    
    for i in range(len(df)):
        progress.progress(
            (i + 1) / len(df),
            text=f"Scoring sentiment: {i + 1} out of {len(df)}"
        )
        news = f"{df.loc[i, 'article_title']} - {df.loc[i, 'article_summary']}"
        text = text_preprocess(news)
        encoded = tokenizer(text, return_tensors="pt", truncation=True, max_length=512)
        with torch.no_grad():
            output = model(**encoded)
        scores = output[0][0].detach().numpy()
        scores = np.exp(scores) / np.sum(np.exp(scores))
        ranking = np.argsort(scores)[::-1]
        curr_pos, curr_neg = 0.0, 0.0
        for r in ranking:
            label = config.id2label[r].lower()
            if "pos" in label:
                curr_pos = scores[r]
            elif "neg" in label:
                curr_neg = scores[r]
        llm_sentiment.append(float(curr_pos - curr_neg))

    progress.empty()

    df["llm_sentiment"] = llm_sentiment
    df["article_date"] = pd.to_datetime(df["article_date"]).dt.date
    df = df.sort_values(by="article_date", ascending=False).reset_index(drop=True)

    cols = ["article_date", "article_title", "article_summary", "article_url",
            "NPS", "compound", "llm_sentiment", "positive_words_list", "negative_words_list"]
    return df[cols]


def risk_bucket(value: float) -> str:
    if value >= 0.5:
        return "green"
    elif value >= 0.25:
        return "amber"
    return "red"


BUCKET_COLOR = {"green": BOB_GREEN, "amber": BOB_AMBER, "red": BOB_RED}

# Shared score-type definitions (column name + display style), used by both the
# per-company time series chart and the cross-company Compare tab.
SCORE_TYPE_OPTIONS = {
    "NPS": "NPS",
    "NLPS (VADER)": "compound",
    "LLM (RoBERTa)": "llm_sentiment",
}
SCORE_TYPE_STYLE = {
    "NPS": {"color": BOB_MAROON, "dash": "solid"},
    "NLPS (VADER)": {"color": BOB_ORANGE, "dash": "dash"},
    "LLM (RoBERTa)": {"color": BOB_RED, "dash": "dot"},
}
# Distinct per-company palette for the Compare tab (BOB colors first, then extras)
COMPARE_PALETTE = [
    BOB_MAROON, BOB_ORANGE, BOB_RED, "#2E86AB", "#5CB85C",
    "#8E44AD", BOB_GOLD, "#E74C3C", "#16A085", "#34495E", "#D4A017",
]


# ----------------------------------------------------------------------------
# Charts (Plotly, for the on-screen interactive dashboard)
# ----------------------------------------------------------------------------
def build_monthly_avg(df: pd.DataFrame) -> pd.DataFrame:
    """Month-wise average of NPS / compound (NLPS) / llm_sentiment for one company."""
    work = df.copy()
    work["article_date"] = pd.to_datetime(work["article_date"])
    work["Year-Month"] = work["article_date"].dt.to_period("M").astype(str)
    return work.groupby("Year-Month")[["NPS", "compound", "llm_sentiment"]].mean().reset_index()


def make_timeseries_chart(df: pd.DataFrame, key_suffix: str):
    """Single-company time series with a multiselect to toggle NPS / NLPS / LLM.

    `key_suffix` must be unique per call site (e.g. derived from the company
    name/tab index) — Streamlit raises a duplicate-element error otherwise
    when this is rendered inside a loop over multiple companies.
    """
    import plotly.graph_objects as go

    avg = build_monthly_avg(df)

    selected_scores = st.multiselect(
        "\n**Score Types**",
        options=list(SCORE_TYPE_OPTIONS.keys()),
        default=list(SCORE_TYPE_OPTIONS.keys()),
        key=f"timeseries_scores_{key_suffix}",
    )

    fig = go.Figure()
    for label in selected_scores:
        column = SCORE_TYPE_OPTIONS[label]
        style = SCORE_TYPE_STYLE[label]
        fig.add_trace(go.Scatter(
            x=avg["Year-Month"], y=avg[column], mode="lines+markers", name=label,
            line=dict(color=style["color"], width=3, dash=style["dash"]),
        ))

    fig.update_layout(
        title="Sentiment Score Time Series",
        template="plotly_white",
        legend_title_text="Score Type",
        margin=dict(l=10, r=10, t=50, b=10),
        height=380,
    )
    return fig


def make_word_bubble_chart(df: pd.DataFrame):
    import plotly.express as px

    translator = str.maketrans("", "", string.punctuation)
    all_words = []
    for col in ["positive_words_list", "negative_words_list"]:
        for entry in df[col]:
            words = entry if isinstance(entry, list) else ast.literal_eval(str(entry))
            for w in words:
                cleaned = w.lower().translate(translator)
                if cleaned:
                    all_words.append(cleaned)

    if not all_words:
        return None

    freq = {}
    for w in all_words:
        if w not in freq:
            freq[w] = [0, 0.0]
        freq[w][0] += 1

    # attach intensity from vader once (cheap, no model needed here)
    from nltk.sentiment import SentimentIntensityAnalyzer
    analyzer = SentimentIntensityAnalyzer()
    for w in freq:
        freq[w][1] = analyzer.polarity_scores(w)["compound"]

    word_df = pd.DataFrame(
        {"Word": list(freq.keys()),
         "Count": [v[0] for v in freq.values()],
         "Intensity": [v[1] for v in freq.values()]}
    )
    word_df["Sentiment"] = np.where(word_df["Intensity"] > 0, "positive", "negative")
    top_pos = word_df[word_df["Sentiment"] == "positive"].nlargest(15, "Count")
    top_neg = word_df[word_df["Sentiment"] == "negative"].nlargest(15, "Count")
    plot_df = pd.concat([top_pos, top_neg])

    if plot_df.empty:
        return None

    fig = px.scatter(
        plot_df, x="Intensity", y="Count", size="Count", color="Sentiment",
        text="Word", color_discrete_map={"positive": BOB_GREEN, "negative": BOB_RED},
        template="plotly_white", title="News Word Intensity vs. Frequency",
    )
    fig.update_traces(textposition="top center")
    fig.update_layout(margin=dict(l=10, r=10, t=50, b=10), height=420)
    return fig


def render_compare_tab(results: dict):
    """Cross-company monthwise comparison — one chart per selected score type,
    one line per selected company, with average score shown in legend.
    Blank/NaN values are ignored; zero values are included in the average.
    """

    st.markdown("**Compare monthly sentiment trends across companies.**")

    company_names = list(results.keys())

    # color_map = {
    #     name: COMPARE_PALETTE[i % len(COMPARE_PALETTE)]
    #     for i, name in enumerate(company_names)
    # }
    
    client = genai.Client(api_key=GEMINI_API_KEY)
    
    prompt = f"""
    Generate exactly {len(company_names)} distinct professional hex color codes,
    one for each company below.
    
    Companies:
    {", ".join(company_names)}
    
    Requirements:
    - Return ONLY the hex color codes
    - Exactly {len(company_names)} colors
    - Comma-separated
    - Format: #RRGGBB
    - Colors must be visually distinct (not in same shade of color, very different than each other)
    - Suitable for a professional financial/consulting dashboard
    - Avoid very light colors that are difficult to see on a white background
    
    Example:
    #1F77B4,#FF7F0E,#2CA02C,#D62728
    """
    
    try:
        response = client.models.generate_content(
            model="gemini-3.5-flash-lite",
            contents=prompt
        )
    
        color_string = response.text.strip()
    
        colors = re.findall(
            r"#[0-9A-Fa-f]{6}",
            color_string
        )
    
        if len(colors) < len(company_names):
            colors = (
                colors + COMPARE_PALETTE
            )[:len(company_names)]
        else:
            colors = colors[:len(company_names)]
    
    except Exception:
        colors = (
            COMPARE_PALETTE
            * ((len(company_names) // len(COMPARE_PALETTE)) + 1)
        )[:len(company_names)]
    
    color_map = {
        name: colors[i]
        for i, name in enumerate(company_names)
    }

    c1, c2 = st.columns(2)

    with c1:
        selected_companies = st.multiselect(
            "Companies",
            options=company_names,
            default=company_names,
            key="compare_companies"
        )

    with c2:
        selected_scores = st.multiselect(
            "Score Types",
            options=list(SCORE_TYPE_OPTIONS.keys()),
            default=list(SCORE_TYPE_OPTIONS.keys()),
            key="compare_score_types",
        )

    if not selected_companies or not selected_scores:
        st.warning("Select at least one company and one score type to compare.")
        return

    # Build monthly averages
    monthly = {
        cname: build_monthly_avg(results[cname]["df"])
        for cname in selected_companies
    }

    import plotly.graph_objects as go

    for label in selected_scores:

        column = SCORE_TYPE_OPTIONS[label]

        fig = go.Figure()

        for cname in selected_companies:

            avg = monthly[cname]

            # --------------------------------------------------
            # Calculate average from original dataframe
            # Ignore only blank / NaN values
            # Keep zero values
            # --------------------------------------------------
            score_value = "N/A"

            if column in results[cname]["df"].columns:

                score_series = pd.to_numeric(
                    results[cname]["df"][column],
                    errors="coerce"
                )

                # mean() ignores NaN automatically
                # Zero values are NOT removed
                if score_series.notna().any():
                    line_avg = score_series.mean()
                    score_value = f"{line_avg:.2f}"

            # --------------------------------------------------
            # Company + average score in legend
            # --------------------------------------------------
            legend_name = f"{cname} [Avg: {score_value}]"

            fig.add_trace(
                go.Scatter(
                    x=avg["Year-Month"],
                    y=avg[column],
                    mode="lines+markers",
                    name=legend_name,
                    line=dict(
                        color=color_map[cname],
                        width=3
                    ),
                )
            )

        fig.update_layout(
            title=f"{label} — Monthwise Comparison",
            template="plotly_white",
            legend_title_text="Company",
            margin=dict(
                l=10,
                r=10,
                t=50,
                b=10
            ),
            height=380,
        )

        st.plotly_chart(
            fig,
            width="stretch",
            key=f"compare_chart_{safe_key(label)}"
        )



# ----------------------------------------------------------------------------
# Excel export (per company) — mirrors the original Dashboard-sheet layout
# ----------------------------------------------------------------------------
def build_excel_bytes(company_name: str, df: pd.DataFrame, start_date, end_date) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import seaborn as sns
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Border, Side
    from openpyxl.formatting.rule import CellIsRule

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.drop(columns=["positive_words_list", "negative_words_list"], errors="ignore").to_excel(
            writer, sheet_name="Data_Coy", index=False
        )
    buf.seek(0)
    wb = load_workbook(buf)

    news_size = len(df)
    avg_nps = df["NPS"].mean()
    avg_nlps = df["compound"].mean()
    avg_llm = df["llm_sentiment"].mean()

    # --- time series chart image ---
    avg = build_monthly_avg(df)

    sns.set(style="whitegrid", context="talk")
    plt.figure(figsize=(12, 8))
    sns.lineplot(x="Year-Month", y="NPS", data=avg, marker="o", label="NPS", linewidth=2.5, color=BOB_MAROON)
    sns.lineplot(x="Year-Month", y="compound", data=avg, marker="o", label="NLPS", linewidth=2.5,
                 linestyle="-.", color=BOB_ORANGE)
    sns.lineplot(x="Year-Month", y="llm_sentiment", data=avg, marker="o", label="LLM", linewidth=2.5,
                 linestyle="--", color=BOB_RED)
    plt.title("Sentiment Score Time Series", fontsize=18, fontweight="bold")
    plt.xticks(rotation=45, fontsize=10)
    plt.tight_layout()
    ts_buf = io.BytesIO()
    plt.savefig(ts_buf, format="png", dpi=100)
    plt.close()
    ts_buf.seek(0)

    # --- bubble chart image ---
    bubble_buf = None
    bubble_fig = make_word_bubble_chart(df)
    
    if bubble_fig is not None:
        translator = str.maketrans("", "", string.punctuation)
        all_words = []
    
        for col in ["positive_words_list", "negative_words_list"]:
            for entry in df[col]:
                words = entry if isinstance(entry, list) else ast.literal_eval(str(entry))
    
                for w in words:
                    cleaned = w.lower().translate(translator)
    
                    if cleaned:
                        all_words.append(cleaned)
    
        # ------------------------------------------------------------
        # Calculate word frequency and sentiment intensity
        # ------------------------------------------------------------
        freq = {}
    
        from nltk.sentiment import SentimentIntensityAnalyzer
        analyzer = SentimentIntensityAnalyzer()
    
        for w in all_words:
            if w not in freq:
                freq[w] = [
                    0,
                    analyzer.polarity_scores(w)["compound"]
                ]
    
            freq[w][0] += 1
    
        # ------------------------------------------------------------
        # Create dataframe
        # ------------------------------------------------------------
        wdf = pd.DataFrame({
            "Keys": list(freq.keys()),
            "Count": [v[0] for v in freq.values()],
            "Intensity": [v[1] for v in freq.values()]
        })
    
        wdf["sentiment"] = np.where(
            wdf["Intensity"] > 0,
            "positive",
            "negative"
        )
    
        # ------------------------------------------------------------
        # Top positive / negative words
        # ------------------------------------------------------------
        top_pos = (
            wdf[wdf["sentiment"] == "positive"]
            .nlargest(15, "Count")
        )
    
        top_neg = (
            wdf[wdf["sentiment"] == "negative"]
            .nlargest(15, "Count")
        )
    
        plot_df = pd.concat(
            [top_neg, top_pos],
            ignore_index=True
        )
    
        # ------------------------------------------------------------
        # Plot
        # ------------------------------------------------------------
        if not plot_df.empty:
    
            fig, ax = plt.subplots(figsize=(14, 8))
    
            sns.scatterplot(
                data=plot_df,
                x="Intensity",
                y="Count",
                size="Count",
                hue="sentiment",
                palette={
                    "positive": BOB_GREEN,
                    "negative": BOB_RED
                },
                sizes=(20, 2000),
                alpha=0.6,
                edgecolor="w",
                linewidth=0.5,
                legend=False,          # IMPORTANT: no seaborn legend
                ax=ax
            )
    
            # --------------------------------------------------------
            # Add word labels
            # --------------------------------------------------------
            for _, row in plot_df.iterrows():
                ax.text(
                    row["Intensity"],
                    row["Count"],
                    row["Keys"],
                    fontsize=9,
                    ha="center",
                    va="center",
                    rotation=25
                )
    
            # --------------------------------------------------------
            # Remove ANY automatically generated legend
            # --------------------------------------------------------
            existing_legend = ax.get_legend()
    
            if existing_legend is not None:
                existing_legend.remove()
    
            # --------------------------------------------------------
            # Create ONLY Positive / Negative legend
            # --------------------------------------------------------
            from matplotlib.lines import Line2D
    
            positive_handle = Line2D(
                [0],
                [0],
                marker="o",
                linestyle="None",
                markerfacecolor=BOB_GREEN,
                markeredgecolor="white",
                markersize=9,
                label="Positive"
            )
    
            negative_handle = Line2D(
                [0],
                [0],
                marker="o",
                linestyle="None",
                markerfacecolor=BOB_RED,
                markeredgecolor="white",
                markersize=9,
                label="Negative"
            )
    
            ax.legend(
                handles=[
                    positive_handle,
                    negative_handle
                ],
                labels=[
                    "Positive",
                    "Negative"
                ],
                title="Sentiment",
                loc="best",
                frameon=True
            )
    
            # --------------------------------------------------------
            # Title
            # --------------------------------------------------------
            ax.set_title(
                "Bubble Plot of Sentiment Analysis",
                fontsize=14
            )
    
            plt.tight_layout()
    
            # --------------------------------------------------------
            # Save image
            # --------------------------------------------------------
            bubble_buf = io.BytesIO()
    
            fig.savefig(
                bubble_buf,
                format="png",
                dpi=100,
                bbox_inches="tight"
            )
    
            plt.close(fig)
    
            bubble_buf.seek(0)
    
    # --- Dashboard sheet ---
    ws = wb.create_sheet("Dashboard")
    for row in ws.iter_rows(min_row=1, max_row=20, max_col=17):
        for cell in row:
            ws.row_dimensions[cell.row].height = 28
    for col_letter in [chr(c) for c in range(ord("A"), ord("Q") + 1)]:
        ws.column_dimensions[col_letter].width = 22

    border = Border(*[Side(border_style="thin", color="000000")] * 4)
    for r in range(3, 16):
        for c in range(1, 18):
            ws.cell(row=r, column=c).border = border

    ws["A1"] = company_name
    ws["A1"].font = Font(color="6E1E33", bold=True, size=20)
    ws["C3"] = "From"
    ws["H3"] = "To"
    ws["D3"] = start_date.strftime("%Y-%m-%d")
    ws["D3"].font = Font(color="F7931E", bold=True, size=16)
    ws["I3"] = end_date.strftime("%Y-%m-%d")
    ws["I3"].font = Font(color="F7931E", bold=True, size=16)
    ws["B4"] = "Number of News Items identified"
    ws["F4"] = news_size
    ws["F4"].font = Font(color="F7931E", bold=True, size=16)

    labels = [("A7", "1-", "B7", "NPS", "F8", avg_nps, "C7",
               "Net percentage score (-1 to 1) = (Positive words - Negative words) / Total sentiment words"),
              ("A10", "2-", "B10", "NLPS", "F11", avg_nlps, "C10",
               "NLTK VADER compound sentiment score, ranges -1 to 1"),
              ("A13", "3-", "B13", "LLM", "F14", avg_llm, "C13",
               "Transformer (RoBERTa) sentiment score, ranges -1 to 1")]
    green = PatternFill(start_color="1FA84C", end_color="1FA84C", fill_type="solid")
    amber = PatternFill(start_color="F4A300", end_color="F4A300", fill_type="solid")
    red = PatternFill(start_color="D93A3A", end_color="D93A3A", fill_type="solid")
    for num_cell, num_txt, lbl_cell, lbl_txt, val_cell, val, desc_cell, desc_txt in labels:
        ws[num_cell] = num_txt
        ws[lbl_cell] = lbl_txt
        ws[lbl_cell].font = Font(bold=True, size=16)
        ws[desc_cell] = desc_txt
        ws[val_cell] = round(float(val), 3) if pd.notna(val) else 0
        ws[val_cell].font = Font(bold=True, size=16)
        ws.conditional_formatting.add(val_cell, CellIsRule(operator="greaterThanOrEqual", formula=["0.5"], fill=green))
        ws.conditional_formatting.add(val_cell, CellIsRule(operator="greaterThanOrEqual", formula=["0.25"], fill=amber))
        ws.conditional_formatting.add(val_cell, CellIsRule(operator="lessThan", formula=["0.25"], fill=red))

    ws["B18"] = "Score Time Series"
    ws["B18"].font = Font(color="6E1E33", bold=True, size=16)
    ws.add_image(XLImage(ts_buf), "B19")
    if bubble_buf is not None:
        ws.add_image(XLImage(bubble_buf), "B60")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------
# Session state
# ----------------------------------------------------------------------------
if "companies" not in st.session_state:
    st.session_state.companies = [{"id": 0}]
if "next_id" not in st.session_state:
    st.session_state.next_id = 1
if "results" not in st.session_state:
    st.session_state.results = {}

# ----------------------------------------------------------------------------
# Header
# ----------------------------------------------------------------------------
st.markdown(
    """
    <div class="bob-banner">
        <h1>🛡️ BOB Reputation Risk Analyzer</h1>
        <p>Media sentiment &amp; reputation risk dashboard — scrape, score, and monitor company news coverage.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

# ----------------------------------------------------------------------------
# Input form
# ----------------------------------------------------------------------------
st.markdown('<div class="section-title">1. Enter Companies</div>', unsafe_allow_html=True)

to_remove = None
for comp in st.session_state.companies:
    cid = comp["id"]
    with st.container(border=True):
        c1, c2, c3, c4, c5, c6 = st.columns([2.2, 1.4, 1.4, 1.6, 1.6, 0.5])
        with c1:
            comp["name"] = st.text_input("Company Name", key=f"name_{cid}", placeholder="e.g. Infosys")
        with c2:
            comp["start_date"] = st.date_input("Start Date", key=f"start_{cid}",
                                                value=date(date.today().year, 1, 1),
                                                min_value=date(1990, 1, 1), max_value=date.today())
        with c3:
            comp["end_date"] = st.date_input("End Date", key=f"end_{cid}",
                                              value=date.today(),
                                              min_value=date(1990, 1, 1), max_value=date.today())
        with c4:
            comp["words"] = st.text_input("Custom Words", key=f"words_{cid}", placeholder="word1|word2")
        with c5:
            comp["scores"] = st.text_input("Custom Scores (-4 to 4)", key=f"scores_{cid}", placeholder="3|-2")
        with c6:
            st.write("")
            st.write("")
            if len(st.session_state.companies) > 1:
                if st.button("✕", key=f"remove_{cid}", help="Remove this company"):
                    to_remove = cid

if to_remove is not None:
    st.session_state.companies = [c for c in st.session_state.companies if c["id"] != to_remove]
    st.rerun()

col_a, col_b = st.columns([1, 5])
with col_a:
    if st.button("+ Add Company"):
        st.session_state.companies.append({"id": st.session_state.next_id})
        st.session_state.next_id += 1
        st.rerun()

run_clicked = st.button("🚀 Run Reputation Risk Analysis", type="primary")

# ----------------------------------------------------------------------------
# Run analysis
# ----------------------------------------------------------------------------
if run_clicked:
    st.session_state.results = {}
        
    vader, tokenizer, config, model = load_engines()

    for comp in st.session_state.companies:
        name = (comp.get("name") or "").strip().upper()
        if not name:
            continue

        start_dt = datetime.combine(comp["start_date"], datetime.min.time())
        end_dt = datetime.combine(comp["end_date"], datetime.min.time())

        if start_dt > end_dt:
            st.error(f"**{name}**: start date is after end date — skipped.")
            continue

        words_list = [w for w in (comp.get("words") or "").split("|") if w.strip()]
        scores_raw = [s for s in (comp.get("scores") or "").split("|") if s.strip()]
        try:
            scores_list = [float(s) for s in scores_raw]
        except ValueError:
            st.error(f"**{name}**: custom scores must be numeric — skipped.")
            continue

        if len(words_list) != len(scores_list):
            st.error(f"**{name}**: custom words count and scores count must match — skipped.")
            continue
        if any(s < -4 or s > 4 for s in scores_list):
            st.error(f"**{name}**: custom scores must be between -4 and 4 — skipped.")
            continue

        custom_words = {w.lower(): s for w, s in zip(words_list, scores_list)}
        search_text = "".join(name.split())

        with st.status(f"Analyzing {name}...", expanded=True) as status:
            log = lambda m: status.write(m)
            links = find_news_links(search_text, log=log)
            if not links:
                status.update(label=f"{name}: no news sources found", state="error")
                continue

            log(f"Found {len(links)} source page(s). Fetching articles between "
                f"{comp['start_date']} and {comp['end_date']}...")
            raw_df = fetch_articles(links, start_dt, end_dt, log=log)
            log(f"Collected {len(raw_df)} raw article(s). Scoring sentiment...")

            if raw_df.empty:
                status.update(label=f"{name}: no articles found in this date range", state="error")
                continue

            scored_df = score_sentiment(raw_df, custom_words, vader, tokenizer, config, model)
            log(f"Sentiment Scoring Completed")
            if scored_df.empty:
                status.update(label=f"{name}: no usable articles after cleanup", state="error")
                continue

            excel_bytes = build_excel_bytes(name, scored_df, comp["start_date"], comp["end_date"])
            st.session_state.results[name] = {
                "df": scored_df,
                "excel": excel_bytes,
                "start": comp["start_date"],
                "end": comp["end_date"],
            }
            status.update(label=f"{name}: analysis complete ✅", state="complete")

            # time.sleep(3)

# ----------------------------------------------------------------------------
# Results (always shown upfront if present in session state)
# ----------------------------------------------------------------------------
if st.session_state.results:
    st.markdown('<div class="section-title">2. Results Dashboard</div>', unsafe_allow_html=True)

    if len(st.session_state.results) > 1:
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w") as zf:
            for cname, res in st.session_state.results.items():
                safe = "".join(c for c in cname if c.isalnum() or c in " _-").strip()
                zf.writestr(f"{safe}.xlsx", res["excel"])
        zip_buf.seek(0)
        st.download_button("⬇️ Download All Reports (ZIP)", data=zip_buf,
                            file_name="reputation_risk_reports.zip", mime="application/zip")

    company_names_list = list(st.session_state.results.keys())
    show_compare = len(company_names_list) > 1
    tab_labels = company_names_list + (["📊 Compare"] if show_compare else [])
    tab_labels = [key.upper() for key in tab_labels]
    tabs = st.tabs(tab_labels)
    company_tabs = tabs[: len(company_names_list)]

    for idx, (tab, cname) in enumerate(zip(company_tabs, company_names_list)):
        res = st.session_state.results[cname]
        with tab:
            df = res["df"]
            # Unique per-tab key fragment — prevents Streamlit's duplicate-element-id
            # error when the same widgets/charts are rendered once per company.
            key_id = f"{idx}_{safe_key(cname)}"

            avg_nps, avg_nlps, avg_llm = df["NPS"].mean(), df["compound"].mean(), df["llm_sentiment"].mean()

            k1, k2, k3, k4 = st.columns(4)
            kpi_defs = [
                ("News Items", f"{len(df)}", BOB_MAROON),
                ("NPS (avg)", f"{avg_nps:.2f}", BUCKET_COLOR[risk_bucket(avg_nps)]),
                ("NLPS — VADER (avg)", f"{avg_nlps:.2f}", BUCKET_COLOR[risk_bucket(avg_nlps)]),
                ("LLM — RoBERTa (avg)", f"{avg_llm:.2f}", BUCKET_COLOR[risk_bucket(avg_llm)]),
            ]
            for col, (label, val, color) in zip([k1, k2, k3, k4], kpi_defs):
                col.markdown(
                    f'<div class="kpi-card" style="background:{color};">'
                    f'<div class="kpi-label">{label}</div>'
                    f'<div class="kpi-value">{val}</div></div>',
                    unsafe_allow_html=True,
                )

            fig = make_timeseries_chart(df, key_suffix=key_id)
            st.plotly_chart(fig, width="stretch", key=f"ts_chart_{key_id}")

            bubble_fig = make_word_bubble_chart(df)
            if bubble_fig is not None:
                st.plotly_chart(bubble_fig, width="stretch", key=f"bubble_chart_{key_id}")

            st.markdown("**Article-level detail**")
            display_df = df.drop(columns=["positive_words_list", "negative_words_list"], errors="ignore")
            display_df.index += 1
            st.dataframe(display_df, width="stretch", height=320)

            st.download_button(
                f"⬇️ Download {cname.upper()} Report (Excel)",
                data=res["excel"],
                file_name=f"{cname.upper()}_reputation_risk_report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"dl_{key_id}",
            )

    if show_compare:
        with tabs[-1]:
            render_compare_tab(st.session_state.results)
else:
    st.info("Add one or more companies above and click **Run Reputation Risk Analysis** — "
            "results, charts and downloads will appear here.")
